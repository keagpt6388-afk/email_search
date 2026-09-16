"""CLI: email-search find / batch / serve."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import typer
from rich.console import Console
from rich.table import Table

from .config import Settings
from .finder import EmailFinder
from .models import Person, SearchResult
from .report import SOURCE_LABEL, to_techgpt, to_xlsx

app = typer.Typer(help="소속·이름·기술분야로 공개 출처에서 이메일을 찾는다.", no_args_is_help=True)
console = Console()


def _print(results: list[SearchResult]) -> None:
    t = Table(title="이메일 검색 결과", show_lines=True)
    for col in ("순위", "전문가", "소속", "이메일", "검증", "출처", "귀속 근거"):
        t.add_column(col, overflow="fold")
    for i, r in enumerate(results, 1):
        b = r.best
        t.add_row(str(i), r.person.name, r.person.affiliation, b.email if b else "-", ("검증" if b.verified else "미검증") if b else "-", SOURCE_LABEL.get(b.source, b.source) if b else "-", b.attribution if b else "-")
    console.print(t)
    for r in results:
        console.print(f"[bold]{r.person.name} / {r.person.affiliation}[/bold]")
        for line in r.log:
            console.print(f"  · {line}")


def _read_people(path: Path, default_field: str) -> list[Person]:
    """xlsx/csv: 열 이름에 '전문가'|'이름'|name, '소속'|affiliation, '기술분야'|field 가 있으면 그것을, 없으면 앞 두 열을 이름·소속으로."""
    rows: list[dict[str, str]] = []
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).active
        header = [str(c.value or "").strip() for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append({header[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row) if i < len(header)})
    else:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            rows = [dict(r) for r in csv.DictReader(fh)]
    people: list[Person] = []
    for r in rows:
        keys = list(r.keys())
        name = next((r[k] for k in keys if k in ("전문가", "이름", "성명", "name", "Name")), r[keys[0]] if keys else "")
        aff = next((r[k] for k in keys if k in ("소속", "기관", "affiliation", "Affiliation", "org")), r[keys[1]] if len(keys) > 1 else "")
        field = next((r[k] for k in keys if k in ("기술분야", "분야", "field", "Field")), "") or default_field
        if name and aff:
            people.append(Person(name=name, affiliation=aff, field=field))
    return people


@app.command()
def find(name: str = typer.Argument(..., help="이름(한글·로마자)"), affiliation: str = typer.Argument(..., help="소속"), field: str = typer.Option("", "--field", "-f", help="기술분야"),
         json_out: Path | None = typer.Option(None, "--json", help="TECH-GPT 응답 표출 규격 JSON 저장 경로"), xlsx: Path | None = typer.Option(None, "--xlsx", help="스프레드시트 저장 경로")) -> None:
    """한 사람 검색."""
    res = EmailFinder(Settings()).find(Person(name=name, affiliation=affiliation, field=field))
    _print([res])
    if json_out:
        json_out.write_text(json.dumps(to_techgpt([res]), ensure_ascii=False, indent=2), encoding="utf-8")
        console.print(f"JSON 저장: {json_out}")
    if xlsx:
        console.print(f"xlsx 저장: {to_xlsx([res], xlsx)}")


@app.command()
def batch(input_path: Path = typer.Argument(..., help="xlsx/csv (열: 전문가·소속·[기술분야])"), field: str = typer.Option("", "--field", "-f", help="기술분야(파일에 열이 없을 때)"),
          out: Path = typer.Option(Path("out/result.xlsx"), "--out", help="xlsx 저장 경로"), json_out: Path | None = typer.Option(None, "--json", help="TECH-GPT 규격 JSON 저장 경로"),
          limit: int = typer.Option(0, "--limit", help="앞에서 N명만")) -> None:
    """파일의 여러 사람을 검색해 xlsx(·JSON)로 저장."""
    people = _read_people(input_path, field)
    if limit:
        people = people[:limit]
    finder = EmailFinder(Settings())
    results: list[SearchResult] = []
    for i, p in enumerate(people, 1):
        console.print(f"[{i}/{len(people)}] {p.name} / {p.affiliation}")
        results.append(finder.find(p))
    _print(results)
    console.print(f"xlsx 저장: {to_xlsx(results, out)}")
    if json_out:
        json_out.write_text(json.dumps(to_techgpt(results), ensure_ascii=False, indent=2), encoding="utf-8")
        console.print(f"JSON 저장: {json_out}")


@app.command()
def serve(host: str = "0.0.0.0", port: int = 8000) -> None:
    """HTTP API (TECH-GPT 도구 호출용)."""
    import uvicorn
    uvicorn.run("email_search.web:app", host=host, port=port)


if __name__ == "__main__":
    app()
